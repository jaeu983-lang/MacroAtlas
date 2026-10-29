"""Tests for raw typed-series XLSX export and freshness revalidation."""
import hashlib
import json
import sqlite3
from datetime import UTC, date, datetime

import pytest
from openpyxl import load_workbook

from macro_data_platform.cli import build_parser, dispatch
from macro_data_platform.config import Settings
from macro_data_platform.contracts import Observation, ObservationStatus, TimeSeries
from macro_data_platform.ingestion.xlsx_export import (
    export_series_task,
    export_typed_series_artifact,
)
from macro_data_platform.storage.history import HistoryStore
from macro_data_platform.storage.typed_series import save_typed_series_artifact


def _artifact(tmp_path, *, observations=None, series_id="FIXTURE.VALUE", country="US",
              unit="USD", dataset_id="fixture-dataset", frequency="annual",
              metadata_evidence=None):
    settings = Settings(tmp_path)
    source_dir = f"fixture-{series_id.replace('.', '-')}-{country.casefold()}"
    raw = settings._inside(f"data/raw/{source_dir}/source.json")
    raw.parent.mkdir(parents=True)
    raw.write_bytes(b'{"source":"fixture"}')
    manifest = settings._inside(f"data/raw/{source_dir}/manifest.json")
    manifest.write_text('{"raw_sha256":"fixture"}', encoding="utf-8")
    series = TimeSeries(
        source_agency="Fixture Agency",
        dataset_id=dataset_id,
        series_id=series_id,
        dimensions={"country": country},
        calendar="gregorian",
        observations=tuple(observations or (
            Observation("2025", 0, ObservationStatus.OBSERVED),
            Observation("2024", None, ObservationStatus.MISSING),
            Observation("2023", "=1+1", ObservationStatus.OBSERVED),
        )),
        retrieved_at=datetime(2026, 10, 28, tzinfo=UTC),
        source_url="https://example.org/fixture",
        raw_ref=f"data/raw/{source_dir}/source.json",
        manifest_ref=f"data/raw/{source_dir}/manifest.json",
        unit=unit,
        frequency=frequency,
        frequency_basis="request_contract" if frequency else "unknown",
    )
    quality = {"status": "passed", "issues": [], "freshness": {
        "status": "FRESH", "policy_id": "frequency_max_age_v2",
    }}
    if metadata_evidence is not None:
        quality["metadata_evidence"] = metadata_evidence
    reference = save_typed_series_artifact(
        settings,
        series=(series,),
        engine_id="fixture/1",
        quality=quality,
    )
    return settings, reference


def _completed_task(settings, reference):
    store = HistoryStore(settings)
    task = HistoryStore.document("comparison fixture", "2026-10-29T00:00:00+00:00")
    task.update({
        "status": "COMPLETED",
        "stages": [{"result": {"payload": {"typed_series_artifact": reference}}}],
    })
    with store.transaction() as connection:
        HistoryStore.insert(connection, task)
    return task["task_id"]


def _comparison_spec(task_a, task_b, *, reviewer_type="AI"):
    return {
        "schema_version": "series-comparison-spec-v1",
        "common_concept_id": "fixture:annual-output",
        "evidence": {
            "reference": "https://example.org/fixture-methodology",
            "sha256": "a" * 64,
            "citation": "Fixture compatibility evidence, section 1",
        },
        "review": {
            "decision": "approved", "reviewer_type": reviewer_type,
            "reviewer_id": "GPT-6-Codex", "reviewed_at": "2026-10-29T00:00:00+00:00",
        },
        "comparison": {
            "date_basis": "typed_period", "aggregation": "none", "transformation": "none",
            "missing_values": "blank", "interpolation": "none", "forward_fill": "none",
            "revision_policy": "as_retrieved",
        },
        "series": [
            {"task_id": task_a, "series_id": "GDP.US", "label": "USA GDP"},
            {"task_id": task_b, "series_id": "GDP.CA", "label": "Canada GDP"},
        ],
    }


def test_comparison_evidence_allows_only_public_json_format_query():
    from macro_data_platform.ingestion.xlsx_export import _validate_comparison_spec

    spec = _comparison_spec("task-us", "task-ca")
    spec["evidence"]["reference"] = (
        "https://api.worldbank.org/v2/indicator/NY.GDP.MKTP.CD?format=json"
    )
    metadata, _ = _validate_comparison_spec(spec)
    assert metadata["evidence_reference"] == spec["evidence"]["reference"]

    for reference in (
        "https://api.worldbank.org/v2/indicator/NY.GDP.MKTP.CD?key=secret",
        "https://api.worldbank.org/v2/indicator/NY.GDP.MKTP.CD?format=json&key=secret",
        "https://api.worldbank.org/v2/indicator/NY.GDP.MKTP.CD?format=xml",
    ):
        spec["evidence"]["reference"] = reference
        with pytest.raises(ValueError, match="comparison_evidence_reference_invalid"):
            _validate_comparison_spec(spec)


def _performance_samples(settings):
    with sqlite3.connect(settings.performance_database_path) as connection:
        connection.row_factory = sqlite3.Row
        return [dict(row) for row in connection.execute(
            "SELECT * FROM performance_samples ORDER BY sample_id"
        )]


def test_raw_xlsx_preserves_values_status_and_hash_linked_provenance(tmp_path):
    settings, reference = _artifact(tmp_path)
    result = export_typed_series_artifact(
        settings, reference, "data/exports/series.xlsx",
        as_of=date(2026, 10, 29),
    )

    output = tmp_path / result["output_path"]
    sidecar = tmp_path / result["provenance_path"]
    workbook = load_workbook(output, data_only=False, read_only=True)
    assert workbook.sheetnames == ["Provenance", "FIXTURE.VALUE"]
    sheet = workbook["FIXTURE.VALUE"]
    headers = [cell.value for cell in sheet[1]]
    value_column = headers.index("raw_value") + 1
    status_column = headers.index("observation_status") + 1
    assert sheet.cell(2, value_column).value == 0
    assert sheet.cell(3, value_column).value is None
    assert sheet.cell(3, status_column).value == "missing"
    formula_like = sheet.cell(4, value_column)
    assert formula_like.value == "=1+1"
    assert formula_like.data_type == "s"
    provenance = workbook["Provenance"]
    provenance_rows = {provenance.cell(row, 1).value: provenance.cell(row, 2).value
                       for row in range(1, provenance.max_row + 1)}
    assert provenance_rows["typed_series_manifest_sha256"] == (
        reference["typed_series_manifest_sha256"]
    )
    manifest = json.loads(sidecar.read_text(encoding="utf-8"))
    assert manifest["output_sha256"] == hashlib.sha256(output.read_bytes()).hexdigest()
    assert manifest["typed_series_sha256"] == reference["typed_series_sha256"]
    assert result["observation_count"] == 3
    [sample] = _performance_samples(settings)
    assert sample["execution_scope"] == "xlsx_export"
    assert sample["engine"] == "typed_xlsx"
    assert sample["stage"] == "export"
    assert sample["status"] == "exported"
    assert sample["sample_class"] == "fixture"
    assert sample["input_rows"] == 3
    assert sample["output_rows"] == 3
    assert sample["input_bytes"] > 0
    assert sample["output_bytes"] > 0
    assert {"quality", "parse", "export"}.issubset(
        json.loads(sample["phase_timings_json"])
    )
    assert not {"task_id", "request_text", "source_url", "raw_value", "credential"} & (
        set(sample)
    )
    serialized_sample = json.dumps(sample, sort_keys=True)
    assert "example.org/fixture" not in serialized_sample
    assert "=1+1" not in serialized_sample
    workbook.close()


def test_raw_xlsx_preserves_ai_reviewed_metadata_mapping_evidence(tmp_path):
    evidence = [{"mapping_id": "fixture-unit-v1", "unit": "USD",
                 "currency": "USD", "reviewer_type": "AI"}]
    settings, reference = _artifact(tmp_path, metadata_evidence=evidence)
    result = export_typed_series_artifact(
        settings, reference, "data/exports/metadata.xlsx", as_of=date(2026, 10, 29),
    )

    workbook = load_workbook(tmp_path / result["output_path"], data_only=True, read_only=True)
    provenance = workbook["Provenance"]
    rows = {provenance.cell(row, 1).value: provenance.cell(row, 2).value
            for row in range(1, provenance.max_row + 1)}
    assert json.loads(rows["metadata_evidence_json"]) == evidence
    sidecar = json.loads((tmp_path / result["provenance_path"]).read_text(encoding="utf-8"))
    assert sidecar["metadata_evidence"] == evidence
    workbook.close()


def test_export_rechecks_freshness_and_creates_no_partial_files(tmp_path):
    settings, reference = _artifact(tmp_path)
    with pytest.raises(ValueError, match="typed_series_not_fresh:STALE"):
        export_typed_series_artifact(
            settings, reference, "data/exports/stale.xlsx",
            as_of=date(2028, 12, 31),
        )
    assert not (tmp_path / "data/exports/stale.xlsx").exists()
    assert not (tmp_path / "data/exports/stale.xlsx.json").exists()


def test_unknown_census_like_frequency_uses_period_format_only_for_freshness(tmp_path):
    settings, reference = _artifact(
        tmp_path, series_id="B01001_001E", country="06", dataset_id="acs/acs1",
        frequency=None,
        observations=(Observation("2025", 10, ObservationStatus.OBSERVED),),
    )
    result = export_typed_series_artifact(
        settings, reference, "data/exports/annual.xlsx", as_of=date(2026, 10, 29),
    )
    assert result["freshness"][0]["frequency"] == "annual"
    assert result["freshness"][0]["frequency_basis"] == "observation_period_format"
    assert result["freshness"][0]["policy_basis"] == "frequency_fallback"


@pytest.mark.parametrize("output", [
    "../outside.xlsx", "/tmp/outside.xlsx", "data/raw/out.xlsx", "data/exports/out.csv",
])
def test_export_rejects_unsafe_or_unsupported_paths(tmp_path, output):
    settings, reference = _artifact(tmp_path)
    with pytest.raises(ValueError, match="output_path_invalid"):
        export_typed_series_artifact(
            settings, reference, output, as_of=date(2026, 10, 29),
        )


def test_export_never_overwrites_existing_workbook_or_sidecar(tmp_path):
    settings, reference = _artifact(tmp_path)
    output = tmp_path / "data/exports/series.xlsx"
    output.parent.mkdir(parents=True)
    output.write_bytes(b"keep")
    with pytest.raises(ValueError, match="output_already_exists"):
        export_typed_series_artifact(
            settings, reference, "data/exports/series.xlsx",
            as_of=date(2026, 10, 29),
        )
    assert output.read_bytes() == b"keep"


def test_export_xlsx_can_resolve_one_artifact_from_completed_task(tmp_path):
    settings, reference = _artifact(tmp_path)
    store = HistoryStore(settings)
    task = HistoryStore.document("fixture task", "2026-10-29T00:00:00+00:00")
    task.update({
        "status": "COMPLETED",
        "stages": [{"result": {"payload": {"typed_series_artifact": reference}}}],
    })
    with store.transaction() as connection:
        HistoryStore.insert(connection, task)

    result = export_series_task(
        settings, task["task_id"], "data/exports/from-task.xlsx",
        as_of=date(2026, 10, 29),
    )
    assert result["task_id"] == task["task_id"]
    assert (tmp_path / result["output_path"]).is_file()
    [sample] = _performance_samples(settings)
    assert sample["engine"] == "typed_xlsx"
    assert "sqlite" in json.loads(sample["phase_timings_json"])


def test_export_xlsx_cli_dispatches_from_task_id(tmp_path):
    (tmp_path / "pyproject.toml").write_text("[project]\nname='fixture'\n", encoding="utf-8")
    settings, reference = _artifact(tmp_path)
    task_id = _completed_task(settings, reference)
    args = build_parser().parse_args([
        "--project-root", str(tmp_path), "export-xlsx", task_id,
        "--output", "data/exports/cli.xlsx",
    ])

    status, result = dispatch(args, clock=lambda: date(2026, 10, 29))
    assert status == 0
    assert result["status"] == "exported"
    assert (tmp_path / result["output_path"]).is_file()


def test_comparison_workbook_preserves_raw_sheets_and_aligns_with_blank_gaps(tmp_path):
    (tmp_path / "pyproject.toml").write_text("[project]\nname='fixture'\n", encoding="utf-8")
    settings, reference_us = _artifact(
        tmp_path, series_id="GDP.US", country="US", observations=(
            Observation("2025", 0, ObservationStatus.OBSERVED),
            Observation("2023", 5, ObservationStatus.OBSERVED),
        ), metadata_evidence=[{"mapping_id": "fixture-unit-v1", "unit": "USD"}],
    )
    _, reference_ca = _artifact(
        tmp_path, series_id="GDP.CA", country="CA", observations=(
            Observation("2025", 7, ObservationStatus.OBSERVED),
            Observation("2024", 8, ObservationStatus.OBSERVED),
        ),
    )
    task_us = _completed_task(settings, reference_us)
    task_ca = _completed_task(settings, reference_ca)
    spec = _comparison_spec(task_us, task_ca)
    spec_path = tmp_path / "data/inbox/comparison.json"
    spec_path.parent.mkdir(parents=True)
    spec_path.write_text(json.dumps(spec), encoding="utf-8")

    args = build_parser().parse_args([
        "--project-root", str(tmp_path), "compare-xlsx", "--spec", "data/inbox/comparison.json",
        "--output", "data/exports/compare.xlsx",
    ])
    status, result = dispatch(args, clock=lambda: date(2026, 10, 29))
    assert status == 0
    workbook = load_workbook(tmp_path / result["output_path"], data_only=False, read_only=True)
    assert workbook.sheetnames == ["Provenance", "GDP.US", "GDP.CA", "Alignment"]
    provenance = workbook["Provenance"]
    provenance_rows = {provenance.cell(row, 1).value: provenance.cell(row, 2).value
                       for row in range(1, provenance.max_row + 1)}
    assert provenance_rows["input_1_typed_series_manifest_sha256"] == (
        reference_us["typed_series_manifest_sha256"]
    )
    assert provenance_rows["input_2_typed_series_manifest_sha256"] == (
        reference_ca["typed_series_manifest_sha256"]
    )
    assert json.loads(provenance_rows["input_1_metadata_evidence_json"]) == [
        {"mapping_id": "fixture-unit-v1", "unit": "USD"},
    ]
    comparison_manifest = json.loads(
        (tmp_path / result["provenance_path"]).read_text(encoding="utf-8")
    )
    assert comparison_manifest["inputs"][0]["metadata_evidence"] == [
        {"mapping_id": "fixture-unit-v1", "unit": "USD"},
    ]
    alignment = workbook["Alignment"]
    header_row = 14
    assert [cell.value for cell in alignment[header_row]] == [
        "typed_period", "USA GDP", "USA GDP__status", "Canada GDP", "Canada GDP__status",
    ]
    assert [alignment.cell(row, 1).value for row in range(header_row + 1, header_row + 4)] == [
        "2023", "2024", "2025",
    ]
    assert alignment.cell(header_row + 1, 2).value == 5
    assert alignment.cell(header_row + 1, 4).value is None
    assert alignment.cell(header_row + 1, 5).value == "no_observation"
    assert alignment.cell(header_row + 2, 2).value is None
    assert alignment.cell(header_row + 2, 3).value == "no_observation"
    assert alignment.cell(header_row + 2, 4).value == 8
    assert alignment.cell(header_row + 3, 2).value == 0
    assert alignment.cell(header_row + 3, 4).value == 7
    manifest = json.loads((tmp_path / result["provenance_path"]).read_text(encoding="utf-8"))
    assert manifest["comparison_review"]["reviewer_type"] == "AI"
    assert manifest["comparison_spec_sha256"] == result["comparison_spec_sha256"]
    assert result["evidence_identity_authenticated"] is False
    [sample] = _performance_samples(settings)
    assert sample["engine"] == "series_compare_xlsx"
    assert sample["sample_class"] == "fixture"
    assert sample["input_rows"] == 4
    assert sample["output_rows"] == 4
    assert {"parse", "sqlite", "quality", "export"}.issubset(
        json.loads(sample["phase_timings_json"])
    )
    serialized_sample = json.dumps(sample, sort_keys=True)
    assert "comparison fixture" not in serialized_sample
    assert "example.org/fixture" not in serialized_sample
    workbook.close()


def test_comparison_fails_closed_for_incompatible_series_and_bad_review(tmp_path):
    from macro_data_platform.ingestion.xlsx_export import export_comparison_spec

    settings, reference_us = _artifact(
        tmp_path, series_id="GDP.US", country="US",
        observations=(Observation("2025", 0, ObservationStatus.OBSERVED),),
    )
    _, reference_ca = _artifact(
        tmp_path, series_id="GDP.CA", country="CA", unit="percent",
        observations=(Observation("2025", 7, ObservationStatus.OBSERVED),),
    )
    task_us, task_ca = _completed_task(settings, reference_us), _completed_task(settings, reference_ca)
    spec = _comparison_spec(task_us, task_ca)
    spec_path = tmp_path / "data/inbox/comparison.json"
    spec_path.parent.mkdir(parents=True)
    spec_path.write_text(json.dumps(spec), encoding="utf-8")

    with pytest.raises(ValueError, match="comparison_series_incompatible"):
        export_comparison_spec(
            settings, "data/inbox/comparison.json", "data/exports/incompatible.xlsx",
            as_of=date(2026, 10, 29),
        )
    assert not (tmp_path / "data/exports/incompatible.xlsx").exists()

    spec["review"]["decision"] = "pending"
    spec_path.write_text(json.dumps(spec), encoding="utf-8")
    with pytest.raises(ValueError, match="comparison_review_not_approved"):
        export_comparison_spec(
            settings, "data/inbox/comparison.json", "data/exports/pending.xlsx",
            as_of=date(2026, 10, 29),
        )
