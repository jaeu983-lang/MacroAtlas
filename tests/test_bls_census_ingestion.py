"""Deterministic BLS/Census connector, safety, and Orchestrator tests."""
import json
import tempfile
import unittest
from datetime import UTC, date, datetime
from functools import partial
from pathlib import Path
from unittest.mock import patch
from urllib.parse import parse_qs, urlsplit

from freshness_support import HISTORICAL_FIXTURE_POLICY
from openpyxl import load_workbook

from macro_data_platform.config import Settings
from macro_data_platform.connectors.bls import BLSQuery, BLSTransport, parse_bls
from macro_data_platform.connectors.census import CensusQuery, CensusTransport, parse_census
from macro_data_platform.connectors.statistical_http import ApiDownload
from macro_data_platform.contracts import EvidenceKind, StageStatus
from macro_data_platform.ingestion.bls import BLSEngine as _BLSEngine
from macro_data_platform.ingestion.census import CensusEngine as _CensusEngine
from macro_data_platform.ingestion.xlsx_export import export_typed_series_artifact
from macro_data_platform.orchestration import build_engine_registry
from macro_data_platform.orchestration.models import StageInput, StageSpec, TaskPlan
from macro_data_platform.storage.history import HistoryStore

BLSEngine = partial(_BLSEngine, freshness_policy=HISTORICAL_FIXTURE_POLICY)
CensusEngine = partial(_CensusEngine, freshness_policy=HISTORICAL_FIXTURE_POLICY)

NOW = datetime(2026, 10, 26, tzinfo=UTC)
BLS_URL = "https://api.bls.gov/publicAPI/v2/timeseries/data/"
CENSUS_URL = "https://api.census.gov/data/2022/acs/acs1"


def download(url, value):
    return ApiDownload(url, json.dumps(value).encode(), 200, NOW, "application/json",
                       EvidenceKind.FIXTURE)


class FixtureTransport:
    def __init__(self, result):
        self.result = result
        self.calls = 0

    def fetch(self, query):
        self.calls += 1
        return self.result


class BLSTests(unittest.TestCase):
    def setUp(self):
        self.query = BLSQuery(("CUUR0000SA0",), 2024, 2024)
        self.body = {
            "status": "REQUEST_SUCCEEDED", "messageCodes": [],
            "Results": {"series": [{"seriesID": "CUUR0000SA0", "data": [{
                "year": "2024", "period": "M01", "periodName": "January",
                "value": "308.417", "footnotes": [],
            }]}]},
        }

    def test_parser_preserves_source_observation_and_unknown_metadata(self):
        result = parse_bls(download(BLS_URL, self.body), self.query)
        self.assertEqual(result["observations"][0]["data"][0]["value"], "308.417")
        self.assertIsNone(result["unit"])
        self.assertIsNone(result["frequency"])

    def test_parser_rejects_wrong_series_and_period(self):
        wrong = json.loads(json.dumps(self.body))
        wrong["Results"]["series"][0]["seriesID"] = "OTHER"
        with self.assertRaises(ValueError):
            BLSQuery(("bad-id/",), 2024, 2024)
        from macro_data_platform.connectors.statistical_http import SourceFailure
        with self.assertRaises(SourceFailure):
            parse_bls(download(BLS_URL, wrong), self.query)
        wrong = json.loads(json.dumps(self.body))
        wrong["Results"]["series"][0]["data"][0]["period"] = "M99"
        with self.assertRaises(SourceFailure):
            parse_bls(download(BLS_URL, wrong), self.query)

    def test_engine_stores_raw_and_manifest_without_credential_fields(self):
        with tempfile.TemporaryDirectory() as directory:
            settings = Settings(Path(directory), network_enabled=True)
            engine = BLSEngine(settings, FixtureTransport(download(BLS_URL, self.body)))
            result = engine.run(StageInput("t", "request", "bls", {
                "series_ids": ["CUUR0000SA0"], "start_year": 2024, "end_year": 2024,
            }, None, ()))
            self.assertEqual(result.status, StageStatus.SUCCEEDED)
            manifest = Path(directory, result.payload["manifest_path"]).read_text()
            self.assertNotIn("registrationkey", manifest.lower())
            self.assertTrue(Path(directory, result.payload["raw_path"]).exists())
            self.assertTrue(Path(
                directory, result.payload["typed_series_artifact"]["typed_series_path"],
            ).is_file())

    def test_engine_exports_catalog_title_and_preserves_unknown_unit_and_currency(self):
        body = json.loads(json.dumps(self.body))
        body["Results"]["series"][0]["catalog"] = {
            "series_title": "All items in U.S. city average",
            "seasonality": "Not Seasonally Adjusted",
            "survey_abbreviation": "CU",
            "survey_name": "Consumer Price Index",
            "measure_data_type": "All items",
            "area": "U.S. city average",
        }
        with tempfile.TemporaryDirectory() as directory:
            settings = Settings(Path(directory), network_enabled=True)
            engine = BLSEngine(settings, FixtureTransport(download(BLS_URL, body)))
            result = engine.run(StageInput("t", "request", "bls", {
                "series_ids": ["CUUR0000SA0"], "start_year": 2024, "end_year": 2024,
            }, None, ()))
            self.assertEqual(result.status, StageStatus.SUCCEEDED)
            evidence, = result.payload["quality"]["metadata_evidence"]
            self.assertEqual(evidence["fields"]["name"]["value"],
                             "All items in U.S. city average")
            self.assertEqual(evidence["fields"]["name"]["source_refs"],
                             ["/Results/series/0/catalog/series_title"])
            self.assertEqual(evidence["fields"]["unit"]["state"], "unknown")
            self.assertIn("series_metadata_partial", result.payload["quality"]["warnings"])
            export = export_typed_series_artifact(
                settings, result.payload["typed_series_artifact"],
                "data/exports/catalog-title.xlsx", as_of=date(2024, 2, 1),
            )
            workbook = load_workbook(
                Path(directory, export["output_path"]), data_only=True, read_only=True,
            )
            sheet = workbook["CUUR0000SA0"]
            headers = [cell.value for cell in next(sheet.iter_rows(min_row=1, max_row=1))]
            row = [cell.value for cell in next(sheet.iter_rows(min_row=2, max_row=2))]
            values = dict(zip(headers, row))
            self.assertEqual(values["series_title"], "All items in U.S. city average")
            self.assertEqual(values["series_title_status"], "known")
            self.assertIsNone(values["unit"])
            self.assertIsNone(values["currency"])
            metadata_sheet = workbook["Series metadata"]
            metadata_headers = [cell.value for cell in next(
                metadata_sheet.iter_rows(min_row=1, max_row=1),
            )]
            metadata_rows = [dict(zip(metadata_headers, row)) for row in metadata_sheet.iter_rows(
                min_row=2, values_only=True,
            )]
            indexed = {row["field"]: row for row in metadata_rows}
            self.assertEqual(indexed["series_title"]["value"], "All items in U.S. city average")
            self.assertEqual(indexed["series_title"]["status"], "known")
            self.assertEqual(indexed["unit"]["value"], "unknown")
            self.assertEqual(indexed["unit"]["status"], "unknown")
            self.assertEqual(indexed["currency"]["value"], "unknown")
            self.assertEqual(indexed["source_attribute:survey_name"]["value"],
                             "Consumer Price Index")
            workbook.close()

    def test_network_disabled_before_transport(self):
        transport = FixtureTransport(download(BLS_URL, self.body))
        engine = BLSEngine(Settings(Path(tempfile.gettempdir())), transport)
        result = engine.run(StageInput("t", "request", "bls", {
            "series_ids": ["CUUR0000SA0"], "start_year": 2024, "end_year": 2024,
        }, None, ()))
        self.assertEqual(result.status, StageStatus.NEEDS_REVIEW)
        self.assertEqual(transport.calls, 0)


class CensusTests(unittest.TestCase):
    def setUp(self):
        self.query = CensusQuery(2022, "acs/acs1", ("NAME", "B01001_001E"), "state:06")
        self.body = [["NAME", "B01001_001E", "state"],
                     ["California", "39029342", "06"]]

    def test_parser_maps_values_without_inventing_unit(self):
        result = parse_census(download(CENSUS_URL, self.body), self.query)
        self.assertEqual(result["observations"][0]["B01001_001E"], "39029342")
        self.assertEqual(result["observations"][0]["state"], "06")
        self.assertIsNone(result["unit"])
        self.assertIsNone(result["frequency"])

    def test_parser_rejects_error_and_shape_changes(self):
        from macro_data_platform.connectors.statistical_http import SourceFailure
        for body in ["Invalid Key", [["NAME"], ["California"]],
                     [["NAME", "OTHER"], ["California", "1"]]]:
            with self.subTest(body=body), self.assertRaises(SourceFailure):
                parse_census(download(CENSUS_URL, body), self.query)

    def test_query_validation_and_engine_manifest_never_contains_key(self):
        for query in [
            lambda: CensusQuery(2022, "../acs", ("NAME",), "state:06"),
            lambda: CensusQuery(2022, "acs/acs1", ("B01001_001E",), "state:06"),
            lambda: CensusQuery(2022, "acs/acs1", ("NAME",), "state:06?key=x"),
            lambda: CensusQuery(2022, "acs/acs1", ("NAME",), "us:*"),
            lambda: CensusQuery(2022, "acs/acs1", ("NAME",), "state:06&key=x"),
        ]:
            with self.subTest(query=query), self.assertRaises(ValueError):
                query()
        with tempfile.TemporaryDirectory() as directory:
            settings = Settings(Path(directory), network_enabled=True)
            engine = CensusEngine(settings, FixtureTransport(download(CENSUS_URL, self.body)))
            result = engine.run(StageInput("t", "request", "census", {
                "year": 2022, "dataset": "acs/acs1", "variables": ["NAME", "B01001_001E"],
                "geography": "state:06",
            }, None, ()))
            self.assertEqual(result.status, StageStatus.SUCCEEDED)
            manifest = json.loads(Path(directory, result.payload["manifest_path"]).read_text())
            serialized = json.dumps(manifest).lower()
            self.assertNotIn('"key"', serialized)
            self.assertNotIn("registrationkey", serialized)
            self.assertNotIn("api.census.gov/data/2022/acs/acs1?", manifest["source_url"])
            self.assertTrue(Path(
                directory, result.payload["typed_series_artifact"]["typed_series_path"],
            ).is_file())

    def test_orchestrator_runs_both_sources(self):
        with tempfile.TemporaryDirectory() as directory:
            settings = Settings(Path(directory), network_enabled=True)
            store = HistoryStore(settings)
            registry = build_engine_registry(
                settings, store,
                bls_transport=FixtureTransport(download(BLS_URL, {
                    "status": "REQUEST_SUCCEEDED", "Results": {"series": [{
                        "seriesID": "CUUR0000SA0", "data": [{
                            "year": "2024", "period": "M01", "value": "308.417",
                        }]
                    }]},
                })),
                census_transport=FixtureTransport(download(CENSUS_URL, self.body)),
                clock=lambda: NOW, freshness_policy=HISTORICAL_FIXTURE_POLICY,
            )
            cases = [
                ("bls", {"series_ids": ["CUUR0000SA0"], "start_year": 2024, "end_year": 2024}),
                ("census", {"year": 2022, "dataset": "acs/acs1",
                            "variables": ["NAME", "B01001_001E"], "geography": "state:06"}),
            ]
            for name, inputs in cases:
                plan = TaskPlan((StageSpec(name, name, inputs),))
                runner = registry.orchestrator_for(plan=plan)
                task = runner.create(name)
                runner.plan(task.task_id, plan)
                result = runner.run(task.task_id)
                self.assertEqual(result["status"], "COMPLETED")


class CredentialTransportTests(unittest.TestCase):
    def test_bls_sends_registered_key_only_to_fixed_official_endpoint(self):
        query = BLSQuery(("CUUR0000SA0",), 2024, 2024)
        with (patch("macro_data_platform.connectors.bls.read_optional_api_key", return_value="fake-key"),
              patch("macro_data_platform.connectors.bls.request_json",
                    return_value=download(BLS_URL, {})) as send):
            result = BLSTransport().fetch(query)
        args, kwargs = send.call_args
        self.assertEqual(args[0], BLS_URL)
        self.assertEqual(kwargs["method"], "POST")
        self.assertEqual(json.loads(kwargs["body"])["registrationkey"], "fake-key")
        self.assertIs(json.loads(kwargs["body"])["catalog"], True)
        self.assertEqual(result.source_url, BLS_URL)

    def test_bls_series_data_does_not_require_an_optional_registration_key(self):
        query = BLSQuery(("CUUR0000SA0",), 2024, 2024)
        with (patch("macro_data_platform.connectors.bls.read_optional_api_key", return_value=None),
              patch("macro_data_platform.connectors.bls.request_json",
                    return_value=download(BLS_URL, {})) as send):
            BLSTransport().fetch(query)
        payload = json.loads(send.call_args.kwargs["body"])
        self.assertNotIn("registrationkey", payload)
        self.assertNotIn("catalog", payload)
        self.assertEqual(payload["seriesid"], ["CUUR0000SA0"])

    def test_census_key_is_removed_from_returned_source_identity(self):
        query = CensusQuery(2022, "acs/acs1", ("NAME", "B01001_001E"), "state:06")
        with (patch("macro_data_platform.connectors.census.read_api_key", return_value="fake-key"),
              patch("macro_data_platform.connectors.census.request_json",
                    return_value=download(CENSUS_URL, [])) as send):
            result = CensusTransport().fetch(query)
        args, kwargs = send.call_args
        parsed = parse_qs(urlsplit(args[0]).query)
        self.assertEqual(kwargs["method"], "GET")
        self.assertEqual(parsed["key"], ["fake-key"])
        self.assertEqual(result.source_url, CENSUS_URL)
        self.assertNotIn("key", urlsplit(result.source_url).query)


if __name__ == "__main__":
    unittest.main()
