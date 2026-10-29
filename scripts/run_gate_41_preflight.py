"""Prepare the offline Gate 41 AI workflow simulation from accepted F/G artifacts."""
import argparse
import hashlib
import json
import sys
from pathlib import Path
from time import perf_counter_ns

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
EVIDENCE = {
    "F": "docs/evidence/gate_f_bls_ai_review_20261029.json",
    "G": "docs/evidence/gate_g_wdi_gdp_ai_review_20261029.json",
}
REQUIRED_FIELDS = {
    "period", "raw_value", "source_agency", "dataset_id", "series_id",
    "dimensions_json", "unit", "currency", "frequency",
}


class PreflightError(ValueError):
    """A recorded artifact cannot safely support the offline simulation."""


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _inside(root: Path, relative: str) -> Path:
    path = (root / relative).resolve()
    try:
        path.relative_to(root.resolve())
    except ValueError:
        raise PreflightError("artifact_path_outside_project") from None
    if not path.is_file():
        raise PreflightError("artifact_missing")
    return path


def _series_summary(sheet, gate: str) -> dict:
    rows = sheet.iter_rows(values_only=True)
    headers = next(rows, None)
    if not headers:
        raise PreflightError("series_sheet_empty")
    columns = {name: index for index, name in enumerate(headers) if isinstance(name, str)}
    if not REQUIRED_FIELDS.issubset(columns):
        raise PreflightError("series_sheet_schema_invalid")

    records = [row for row in rows if any(value is not None for value in row)]
    periods = [str(row[columns["period"]]) for row in records]
    if not records or any(not period for period in periods):
        raise PreflightError("series_observations_missing")

    first = records[0]
    dimensions = json.loads(first[columns["dimensions_json"]] or "{}")
    if not isinstance(dimensions, dict):
        raise PreflightError("series_dimensions_invalid")
    summary = {
        "series_id": first[columns["series_id"]],
        "publisher": first[columns["source_agency"]],
        "dataset_id": first[columns["dataset_id"]],
        "unit": first[columns["unit"]],
        "currency": first[columns["currency"]],
        "frequency": first[columns["frequency"]],
        "dimension_names": sorted(dimensions),
        "period_start": min(periods),
        "period_end": max(periods),
        "observation_count": len(records),
        "missing_count": sum(row[columns["raw_value"]] is None for row in records),
        "zero_count": sum(row[columns["raw_value"]] in (0, "0", "0.0") for row in records),
    }
    if gate == "G":
        summary["country"] = dimensions.get("country_iso3")
    return summary


def inspect_workbook(path: Path, gate: str, expected_sha256: str,
                     sidecar: dict) -> dict:
    actual_sha256 = _sha256(path)
    if actual_sha256 != expected_sha256 or sidecar.get("output_sha256") != expected_sha256:
        raise PreflightError("workbook_hash_mismatch")

    workbook = load_workbook(path, read_only=True, data_only=False)
    try:
        names = workbook.sheetnames
        if "Provenance" not in names:
            raise PreflightError("provenance_sheet_missing")
        if gate == "G" and "Alignment" not in names:
            raise PreflightError("alignment_sheet_missing")

        formulas = sum(
            cell.data_type == "f"
            for sheet in workbook.worksheets
            for row in sheet.iter_rows()
            for cell in row
        )
        if formulas:
            raise PreflightError("formula_cells_present")

        data_sheets = [
            sheet for sheet in workbook.worksheets
            if sheet.title not in {"Provenance", "Alignment"}
        ]
        series = [_series_summary(sheet, gate) for sheet in data_sheets]
        if len(series) != sidecar.get("series_count"):
            raise PreflightError("series_count_mismatch")
        if sum(item["observation_count"] for item in series) != sidecar.get(
                "observation_count"):
            raise PreflightError("observation_count_mismatch")

        aligned_periods = []
        if gate == "G":
            alignment = workbook["Alignment"]
            values = list(alignment.iter_rows(min_col=1, max_col=1, values_only=True))
            header_index = next(
                (index for index, row in enumerate(values) if row[0] == "typed_period"), None,
            )
            if header_index is None:
                raise PreflightError("alignment_period_header_missing")
            aligned_periods = [
                str(row[0]) for row in values[header_index + 1:] if row[0] is not None
            ]
            common_periods = set.intersection(*[
                {period for period in (item["period_start"], item["period_end"])}
                for item in series
            ])
            if len(series) != 2 or len(aligned_periods) != 1 or set(aligned_periods) != common_periods:
                raise PreflightError("comparison_alignment_mismatch")
    finally:
        workbook.close()

    freshness = sidecar.get("freshness", [])
    states = [item.get("status") for item in freshness if isinstance(item, dict)]
    if not states or any(state != "FRESH" for state in states):
        raise PreflightError("artifact_not_fresh")

    return {
        "status": "READY",
        "workbook_sha256": actual_sha256,
        "sheet_names": names,
        "series": series,
        "freshness": states,
        "formula_cells": formulas,
        "aligned_periods": aligned_periods,
    }


def _gate_input(root: Path, gate: str) -> dict:
    evidence = json.loads(_inside(root, EVIDENCE[gate]).read_text(encoding="utf-8"))
    if evidence.get("decision") != "accepted_scoped" or evidence.get("reviewer_type") != "AI":
        raise PreflightError(f"gate_{gate.lower()}_acceptance_record_invalid")

    if gate == "F":
        task = evidence["live_task"]
        output_path, output_sha = task["workbook_path"], task["workbook_sha256"]
        sidecar_path, sidecar_sha = task["provenance_path"], task["provenance_sha256"]
    else:
        output = evidence["output"]
        output_path, output_sha = output["path"], output["sha256"]
        sidecar_path, sidecar_sha = output["provenance_path"], output["provenance_sha256"]

    workbook_path = _inside(root, output_path)
    sidecar_file = _inside(root, sidecar_path)
    if _sha256(sidecar_file) != sidecar_sha:
        raise PreflightError("provenance_hash_mismatch")
    sidecar = json.loads(sidecar_file.read_text(encoding="utf-8"))
    result = inspect_workbook(workbook_path, gate, output_sha, sidecar)
    result["provenance_sha256"] = sidecar_sha
    return result


def build_report(root: Path) -> dict:
    started = perf_counter_ns()
    gates = {}
    for gate in ("F", "G"):
        gate_started = perf_counter_ns()
        gates[gate] = _gate_input(root, gate)
        gates[gate]["preflight_ms"] = round((perf_counter_ns() - gate_started) / 1_000_000, 3)
    return {
        "status": "READY",
        "network_requested": False,
        "database_accessed": False,
        "files_written": False,
        "raw_values_emitted": False,
        "gates": gates,
        "simulation_tasks": [
            "BLS: identify publisher, series, frequency, period range, freshness, and missing-versus-zero semantics; report blank dataset/unit/currency as unknown.",
            "WDI: identify countries, shared indicator, unit, period, alignment policy, and interpretation limits.",
        ],
        "feedback_rule": "record only a real correction; label any generated record AI-simulated",
        "total_preflight_ms": round((perf_counter_ns() - started) / 1_000_000, 3),
    }


def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--project-root", type=Path, default=ROOT)
    args = parser.parse_args(argv)
    try:
        report = build_report(args.project_root.resolve())
    except (OSError, KeyError, TypeError, ValueError, json.JSONDecodeError) as exc:
        reason = exc.args[0] if isinstance(exc, PreflightError) else "input_validation_failed"
        report = {
            "status": "BLOCKED",
            "network_requested": False,
            "database_accessed": False,
            "files_written": False,
            "raw_values_emitted": False,
            "reason": reason,
        }
    print(json.dumps(report, ensure_ascii=False, sort_keys=True))
    return 0 if report["status"] == "READY" else 2


if __name__ == "__main__":
    sys.exit(main())
