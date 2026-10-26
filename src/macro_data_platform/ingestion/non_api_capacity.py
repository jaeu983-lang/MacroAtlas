"""Bounded non-API inspection and durable, deduplicated deferral records."""
import csv
import hashlib
import io
import json
import sqlite3
import zipfile
from contextlib import contextmanager
from datetime import UTC, datetime
from pathlib import PurePosixPath

from macro_data_platform.connectors.tabular_files import (
    MAX_CELLS,
    MAX_COLUMNS,
    MAX_FIELD_CHARS,
    MAX_INPUT_BYTES,
    MAX_ROWS,
    MAX_ZIP_ENTRIES,
    MAX_ZIP_RATIO,
    MAX_ZIP_UNCOMPRESSED_BYTES,
)

POLICY_VERSION = 'non_api_v1'


def decision(reason, value=None, threshold=None, status='deferred'):
    return {'reason': reason, 'observed': value, 'threshold': threshold, 'status': status,
            'policy_version': POLICY_VERSION,
            'unit': {'input_bytes': 'bytes', 'member_bytes': 'bytes',
                     'download_bytes': 'bytes', 'zip_uncompressed_bytes': 'bytes',
                     'zip_entries': 'entries', 'zip_ratio': 'ratio', 'rows': 'rows',
                     'columns': 'columns', 'cells': 'cells', 'field_chars': 'characters'
                     }.get(reason),
            'inspection_point': 'bounded_preflight'}


def inspect_file(path, format_name, options):
    size = path.stat().st_size
    if size >= MAX_INPUT_BYTES:
        return decision('input_bytes', size, MAX_INPUT_BYTES)
    with path.open('rb') as stream:
        raw = stream.read(MAX_INPUT_BYTES)
    if len(raw) >= MAX_INPUT_BYTES:
        return decision('input_bytes', len(raw), MAX_INPUT_BYTES)
    return inspect_bytes(raw, format_name, options)


def inspect_bytes(raw, format_name, options):
    if format_name in {'zip', 'xlsx'}:
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                entries = archive.infolist()
                if len(entries) >= MAX_ZIP_ENTRIES:
                    return decision('zip_entries', len(entries), MAX_ZIP_ENTRIES)
                total = sum(item.file_size for item in entries)
                if total >= MAX_ZIP_UNCOMPRESSED_BYTES:
                    return decision('zip_uncompressed_bytes', total, MAX_ZIP_UNCOMPRESSED_BYTES)
                for item in entries:
                    name = item.filename
                    if (item.flag_bits & 1 or '\\' in name or name.startswith('/')
                            or '..' in PurePosixPath(name).parts or ':' in name
                            or ((item.external_attr >> 16) & 0o170000) == 0o120000):
                        return decision('unsafe_archive', status='security_review')
                    ratio = item.file_size / max(item.compress_size, 1)
                    if ratio >= MAX_ZIP_RATIO:
                        return decision('zip_ratio', ratio, MAX_ZIP_RATIO, 'security_review')
                if format_name == 'zip':
                    members = [item for item in entries if not item.is_dir()
                               and PurePosixPath(item.filename).suffix.lower() in {
                                   '.csv', '.tsv', '.txt', '.html', '.htm', '.xls', '.xlsx'}]
                    selected = options.get('zip_member')
                    if selected:
                        members = [item for item in members if item.filename == selected]
                    if len(members) != 1:
                        return decision('archive_selection', status='inspection_required')
                    item = members[0]
                    if item.file_size >= MAX_INPUT_BYTES:
                        return decision('member_bytes', item.file_size, MAX_INPUT_BYTES)
                    with archive.open(item) as stream:
                        body = stream.read(MAX_INPUT_BYTES)
                    if len(body) >= MAX_INPUT_BYTES:
                        return decision('member_bytes', len(body), MAX_INPUT_BYTES)
                    extension = PurePosixPath(item.filename).suffix.lstrip('.').lower()
                    extension = 'html' if extension == 'htm' else extension
                    if extension == 'zip':
                        return decision('nested_archive', status='inspection_required')
                    return inspect_bytes(body, extension, options)
        except (OSError, ValueError, zipfile.BadZipFile, RuntimeError):
            return decision('invalid_archive', status='security_review')
    if format_name in {'csv', 'tsv', 'txt'}:
        if format_name == 'txt' and options.get('delimiter') is None:
            return decision('delimiter_unknown', status='inspection_required')
        try:
            encoding = 'utf-16' if raw.startswith((b'\xff\xfe', b'\xfe\xff')) else 'utf-8-sig'
            delimiter = options.get('delimiter', '\t' if format_name == 'tsv' else ',')
            if not isinstance(delimiter, str) or len(delimiter) != 1:
                return decision('delimiter_unknown', status='inspection_required')
            previous = csv.field_size_limit()
            try:
                csv.field_size_limit(MAX_FIELD_CHARS)
                reader = csv.reader(io.StringIO(raw.decode(encoding)), delimiter=delimiter)
                cells = 0
                header = options.get('header_row', 1)
                if type(header) is not int or header < 1:
                    return decision('header_unknown', status='inspection_required')
                for number, row in enumerate(reader, 1):
                    cells += len(row)
                    for reason, value, threshold in (
                        ('rows', max(0, number - header), MAX_ROWS),
                        ('columns', len(row), MAX_COLUMNS), ('cells', cells, MAX_CELLS),
                        ('field_chars', max(map(len, row), default=0), MAX_FIELD_CHARS),
                    ):
                        if value >= threshold:
                            return decision(reason, value, threshold)
            finally:
                csv.field_size_limit(previous)
        except (UnicodeError, csv.Error, TypeError, ValueError):
            return decision('text_inspection_incomplete', status='inspection_required')
    if format_name in {'xls', 'xlsx'}:
        try:
            if format_name == 'xlsx':
                import openpyxl
                book = openpyxl.load_workbook(io.BytesIO(raw), read_only=True,
                                               data_only=False, keep_links=False)
                try:
                    name = options.get('sheet_name')
                    if name is None and len(book.sheetnames) != 1:
                        return decision('worksheet_selection', status='inspection_required')
                    sheet = book[name] if name else book[book.sheetnames[0]]
                    rows, columns = sheet.max_row, sheet.max_column
                finally:
                    book.close()
            else:
                import xlrd
                book = xlrd.open_workbook(file_contents=raw, on_demand=True)
                try:
                    name = options.get('sheet_name')
                    if name is None and book.nsheets != 1:
                        return decision('worksheet_selection', status='inspection_required')
                    sheet = book.sheet_by_name(name) if name else book.sheet_by_index(0)
                    rows, columns = sheet.nrows, sheet.ncols
                finally:
                    book.release_resources()
            if rows is None or columns is None:
                return decision('worksheet_dimensions_unknown', status='inspection_required')
            for reason, value, threshold in (
                ('rows', max(0, rows - options.get('header_row', 1)), MAX_ROWS),
                ('columns', columns, MAX_COLUMNS), ('cells', rows * columns, MAX_CELLS),
            ):
                if value >= threshold:
                    return decision(reason, value, threshold)
        except Exception:  # noqa: BLE001 - parser diagnostics must not escape into task output
            return decision('worksheet_inspection_incomplete', status='inspection_required')
    return None


class DeferredNonAPIStore:
    def __init__(self, settings, *, read_only=False):
        self.settings = settings
        self.read_only = read_only
        if read_only:
            return
        settings.database_path.parent.mkdir(parents=True, exist_ok=True)
        with self.connect() as db:
            db.execute('CREATE TABLE IF NOT EXISTS deferred_non_api '
                       '(id TEXT PRIMARY KEY, document TEXT NOT NULL)')
            db.execute('CREATE TABLE IF NOT EXISTS deferred_non_api_events '
                       '(seq INTEGER PRIMARY KEY, id TEXT NOT NULL, document TEXT NOT NULL)')

    @contextmanager
    def connect(self):
        db = (sqlite3.connect(self.settings.database_path.as_uri() + "?mode=ro", uri=True)
              if self.read_only else sqlite3.connect(self.settings.database_path))
        try:
            with db:
                yield db
        finally:
            db.close()

    def register(self, scope, task_id, stage_id, result):
        if self.read_only:
            raise ValueError("deferral store is read-only")
        safe_scope = {key: scope[key] for key in (
            'relative_path', 'source_url', 'source_agency', 'source_table_name', 'format',
            'delimiter', 'header_row', 'zip_member', 'sheet_name', 'table_index',
        ) if key in scope}
        from urllib.parse import urlsplit
        url = safe_scope.get('source_url')
        if url:
            parsed = urlsplit(url)
            if (parsed.scheme != 'https' or not parsed.hostname or parsed.username
                    or parsed.password or parsed.query or parsed.fragment
                    or len(url) > 512 or any(char.isspace() for char in url)):
                raise ValueError('deferral URL must be credential-free HTTPS')
        identity = {key: value for key, value in safe_scope.items()
                    if key != 'relative_path' or not safe_scope.get('source_url')}
        key = hashlib.sha256(json.dumps([identity, POLICY_VERSION], sort_keys=True).encode()).hexdigest()
        document = {'id': key, 'scope': safe_scope, 'task_id': task_id, 'stage_id': stage_id,
                    'observed_at': datetime.now(UTC).isoformat(), **result}
        encoded = json.dumps(document, ensure_ascii=False, allow_nan=False)
        with self.connect() as db:
            db.execute('INSERT INTO deferred_non_api VALUES (?, ?) ON CONFLICT(id) '
                       'DO UPDATE SET document=excluded.document', (key, encoded))
            db.execute('INSERT INTO deferred_non_api_events(id,document) VALUES (?,?)', (key, encoded))
        return document

    def list(self, record_id=None):
        if not self.settings.database_path.exists():
            return []
        with self.connect() as db:
            if not db.execute("SELECT 1 FROM sqlite_master WHERE type='table' "
                              "AND name='deferred_non_api'").fetchone():
                return []
            if record_id:
                rows = db.execute('SELECT document FROM deferred_non_api_events WHERE id=? '
                                  'ORDER BY seq DESC LIMIT 100', (record_id,)).fetchall()
            else:
                rows = db.execute('SELECT document FROM deferred_non_api ORDER BY '
                                  '(SELECT MAX(seq) FROM deferred_non_api_events e '
                                  'WHERE e.id=deferred_non_api.id) DESC LIMIT 100').fetchall()
        return [json.loads(row[0]) for row in rows]
