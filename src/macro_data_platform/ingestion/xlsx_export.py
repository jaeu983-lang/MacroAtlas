"""Write-once XLSX export from integrity-checked, freshness-passed typed artifacts."""

import hashlib
import io
import json
import logging
import os
import re
import tempfile
from collections.abc import Mapping
from datetime import UTC, date, datetime
from functools import wraps
from pathlib import Path
from time import perf_counter_ns
from urllib.parse import urlsplit

from openpyxl import Workbook, load_workbook

from macro_data_platform.config import Settings
from macro_data_platform.ingestion.freshness import (
    DEFAULT_FRESHNESS_POLICY,
    FreshnessPolicy,
    FreshnessStatus,
    assess_freshness,
)
from macro_data_platform.performance.metrics import (
    PerformanceMetrics,
    current_metrics,
    measure_current,
    runtime_metadata,
)
from macro_data_platform.performance.storage import append_sample, build_sample
from macro_data_platform.storage.history import HistoryStore
from macro_data_platform.storage.typed_series import load_typed_series_artifact

MAX_XLSX_BYTES = 64 * 1024 * 1024
MAX_XLSX_OBSERVATIONS = 1_048_575
MAX_XLSX_COLUMNS = 16_384
_PERIOD_FORMAT = {
    "daily": "YYYY-MM-DD",
    "weekly": "YYYY-Www",
    "monthly": "YYYY-MM",
    "quarterly": "YYYY-Qn",
    "semiannual": "YYYY-Hn",
    "annual": "YYYY",
}
_XLSX_INVALID_SHEET_CHARS = re.compile(r"[\\/*?:\[\]]")
_LOGGER = logging.getLogger(__name__)


def _source_sample_class(engine_id: object) -> str:
    if isinstance(engine_id, str) and engine_id.startswith("fixture/"):
        return "fixture"
    if isinstance(engine_id, str) and engine_id.startswith("synthetic/"):
        return "synthetic"
    return "stored"


def _measured_xlsx_export(engine: str):
    """Record only labels, counts, bytes, phases, and elapsed time for exports."""
    def decorate(function):
        @wraps(function)
        def wrapped(settings: Settings, *args, **kwargs):
            if current_metrics() is not None:
                return function(settings, *args, **kwargs)
            metrics = PerformanceMetrics()
            metrics.observe_sample_class("stored")
            started_ns = perf_counter_ns()
            status = "FAILED"
            try:
                with metrics.activate():
                    result = function(settings, *args, **kwargs)
                status = result.get("status", "UNKNOWN") if isinstance(result, dict) else "UNKNOWN"
                return result
            finally:
                try:
                    append_sample(settings, build_sample(
                        observed_at=datetime.now(UTC).isoformat(),
                        execution_scope="xlsx_export",
                        engine=engine,
                        engine_version="1",
                        stage="export",
                        status=status,
                        duration_ns=max(0, perf_counter_ns() - started_ns),
                        metrics=metrics.snapshot(),
                        environment=runtime_metadata(str(settings.project_root)),
                    ))
                except Exception:  # noqa: BLE001 - telemetry cannot change export outcome
                    _LOGGER.warning("XLSX performance telemetry sample could not be recorded")
        return wrapped
    return decorate


def _artifact_input_bytes(settings: Settings, reference: Mapping, document: Mapping) -> int:
    total = 0
    for key in ("typed_series_path", "typed_series_manifest_path"):
        total += settings._inside(reference[key]).stat().st_size
    total += sum(
        row["bytes"] for row in document.get("source_refs", [])
        if isinstance(row, Mapping) and type(row.get("bytes")) is int
        and row["bytes"] >= 0
    )
    return total


def _record_export_shape(
    document: Mapping,
    *,
    input_bytes: int,
    output_bytes: int | None = None,
    observation_count: int = 0,
    series_count: int | None = None,
) -> None:
    metrics = current_metrics()
    if metrics is None:
        return
    metrics.observe_sample_class(_source_sample_class(document.get("engine_id")))
    metrics.observe_input(
        byte_count=input_bytes,
        rows=observation_count,
        columns=4 if observation_count else 0,
    )
    if output_bytes is not None:
        metrics.observe_output(
            byte_count=output_bytes,
            rows=observation_count,
            columns=max(20, 1 + 2 * (series_count or len(document.get("series", [])))),
        )


def _period_format(contract: Mapping) -> str | None:
    frequency_format = _PERIOD_FORMAT.get(contract.get("frequency"))
    if frequency_format:
        return frequency_format
    periods = [row.get("period") for row in contract.get("observations", [])
               if isinstance(row, Mapping) and isinstance(row.get("period"), str)]
    if not periods:
        return None
    patterns = (
        (re.compile(r"^\d{4}$"), "YYYY"),
        (re.compile(r"^\d{4}-\d{2}$"), "YYYY-MM"),
        (re.compile(r"^\d{4}-\d{2}-\d{2}$"), "YYYY-MM-DD"),
        (re.compile(r"^\d{4}-W\d{2}$"), "YYYY-Www"),
        (re.compile(r"^\d{4}-Q[1-4]$"), "YYYY-Qn"),
        (re.compile(r"^\d{4}-H[1-2]$"), "YYYY-Hn"),
    )
    for pattern, label in patterns:
        if all(pattern.fullmatch(period) for period in periods):
            return label
    return None


def _series_key(contract: Mapping) -> dict[str, str] | None:
    agency = contract.get("source_agency")
    dataset = contract.get("dataset_id")
    series_id = contract.get("series_id")
    dimensions = contract.get("dimensions")
    if not isinstance(dimensions, Mapping):
        return None
    if agency == "U.S. Census Bureau" and dataset and series_id:
        geography = dimensions.get("geography")
        if geography is None and len(dimensions) == 1:
            dimension_name, dimension_value = next(iter(dimensions.items()))
            if isinstance(dimension_name, str) and isinstance(dimension_value, str):
                geography = f"{dimension_name}:{dimension_value}"
        if isinstance(geography, str):
            return {"dataset": dataset, "variable": series_id, "geography": geography}
    if agency == "Bank of Japan" and dataset and series_id:
        return {"database": dataset, "series_code": series_id}
    if agency == "World Bank" and dataset and series_id:
        country = dimensions.get("country_iso3")
        if isinstance(country, str):
            return {"country": country, "indicator_id": series_id, "source": dataset}
    if agency and series_id:
        key = {"source_agency": agency, "series_id": series_id}
        if isinstance(dataset, str):
            key["dataset_id"] = dataset
        key.update({f"dimension:{name}": value for name, value in dimensions.items()
                    if isinstance(name, str) and isinstance(value, str)})
        return key
    return None


def _reevaluate_freshness(document: Mapping, policy: FreshnessPolicy,
                          as_of: datetime | date) -> list[dict]:
    results = []
    for entry in document["series"]:
        contract = entry["contract"]
        observations = contract["observations"]
        periods = [row["period"] for row in observations
                   if row["status"] == "observed" and row["raw_value"] is not None]
        result = assess_freshness(
            periods,
            frequency=contract.get("frequency"),
            frequency_basis=contract.get("frequency_basis"),
            period_format=_period_format(contract),
            series_id=contract.get("series_id"),
            series_key=_series_key(contract),
            as_of=as_of,
            policy=policy,
        )
        if result["status"] != FreshnessStatus.FRESH.value:
            raise ValueError(f"typed_series_not_fresh:{result['status']}:{result['reason']}")
        results.append(result)
    return results


def _sheet_name(series_id: str, used: set[str]) -> str:
    cleaned = _XLSX_INVALID_SHEET_CHARS.sub("_", series_id).strip("' ") or "Series"
    if cleaned.casefold() == "provenance":
        cleaned = "Series_Provenance"
    base = cleaned[:31]
    candidate, index = base, 1
    while candidate.casefold() in used:
        suffix = f"_{index}"
        candidate = f"{base[:31 - len(suffix)]}{suffix}"
        index += 1
    used.add(candidate.casefold())
    return candidate


def _set_literal(cell, value) -> None:
    """Store source strings as literal text, including formula-looking values."""
    cell.value = value
    if isinstance(value, str):
        cell.data_type = "s"


def _source_hashes(document: Mapping) -> dict[tuple[str, str], dict]:
    return {
        (row["role"], row["path"]): row
        for row in document.get("source_refs", [])
        if isinstance(row, dict) and isinstance(row.get("role"), str)
        and isinstance(row.get("path"), str)
    }


def _build_raw_workbook(
    document: Mapping,
    freshness: list[dict],
    *,
    comparison_metadata: Mapping | None = None,
) -> tuple[bytes, int]:
    workbook = Workbook()
    provenance = workbook.active
    provenance.title = "Provenance"
    provenance_rows = [
        ("field", "value"),
        ("schema_version", document["schema_version"]),
        ("engine_id", document["engine_id"]),
        ("artifact_created_at", document["created_at"]),
        ("artifact_sha256", document.get("_artifact_sha256", "multiple_inputs")),
        ("typed_series_manifest_path", document.get("_typed_series_manifest_path")),
        ("typed_series_manifest_sha256", document.get("_typed_series_manifest_sha256")),
        ("freshness_rechecked_at", freshness[0]["evaluated_at"]),
        ("series_count", len(document["series"])),
    ]
    for row_index, row in enumerate(provenance_rows, start=1):
        for column_index, value in enumerate(row, start=1):
            _set_literal(provenance.cell(row_index, column_index), value)
    for index, result in enumerate(freshness, start=1):
        metadata = (
            f"series_{index}_freshness_status", result["status"],
            f"series_{index}_freshness_policy", result["policy_id"],
            f"series_{index}_freshness_policy_hash", result["policy_fingerprint_sha256"],
            f"series_{index}_freshness_basis", result["policy_basis"],
            f"series_{index}_latest_period", result["latest_period"],
            f"series_{index}_age_days", result["age_days"],
            f"series_{index}_max_age_days", result["max_age_days"],
        )
        for offset in range(0, len(metadata), 2):
            row = provenance.max_row + 1
            _set_literal(provenance.cell(row, 1), metadata[offset])
            _set_literal(provenance.cell(row, 2), metadata[offset + 1])
    for index, artifact in enumerate(document.get("_input_artifacts", []), start=1):
        for key, value in artifact.items():
            row = provenance.max_row + 1
            _set_literal(provenance.cell(row, 1), f"input_{index}_{key}")
            _set_literal(provenance.cell(row, 2), value)
    if comparison_metadata:
        for key, value in comparison_metadata.items():
            row = provenance.max_row + 1
            _set_literal(provenance.cell(row, 1), f"comparison_{key}")
            _set_literal(provenance.cell(row, 2), value)

    hashes = _source_hashes(document)
    used = {"provenance"}
    observation_count = 0
    headers = (
        "period", "source_period", "raw_value", "observation_status",
        "source_agency", "dataset_id", "series_id", "dimensions_json", "calendar",
        "unit", "currency", "frequency", "frequency_basis", "source_revision",
        "retrieved_at", "source_url", "raw_path", "raw_sha256", "manifest_path",
        "manifest_sha256",
    )
    if len(headers) > MAX_XLSX_COLUMNS:
        raise ValueError("xlsx_column_limit")
    for index, entry in enumerate(document["series"]):
        contract = entry["contract"]
        sheet = workbook.create_sheet(_sheet_name(contract["series_id"], used))
        for column, value in enumerate(headers, start=1):
            _set_literal(sheet.cell(1, column), value)
        raw_ref, manifest_ref = contract["raw_ref"], contract["manifest_ref"]
        raw_hash = hashes.get(("raw", raw_ref), {}).get("sha256")
        manifest_hash = hashes.get(("manifest", manifest_ref), {}).get("sha256")
        dimensions = json.dumps(contract["dimensions"], ensure_ascii=False, sort_keys=True)
        row_count = len(contract["observations"])
        if row_count > MAX_XLSX_OBSERVATIONS:
            raise ValueError("xlsx_observation_limit")
        for observation_index, observation in enumerate(contract["observations"], start=2):
            values = (
                observation["period"], observation["source_period"],
                observation["raw_value"], observation["status"], contract["source_agency"],
                contract["dataset_id"], contract["series_id"], dimensions, contract["calendar"],
                contract["unit"], contract["currency"], contract["frequency"],
                contract["frequency_basis"], contract["source_revision"],
                contract["retrieved_at"], contract["source_url"], raw_ref, raw_hash,
                manifest_ref, manifest_hash,
            )
            for column, value in enumerate(values, start=1):
                _set_literal(sheet.cell(observation_index, column), value)
        observation_count += row_count
        _set_literal(provenance.cell(provenance.max_row + 1, 1), f"series_{index + 1}_sheet")
        _set_literal(provenance.cell(provenance.max_row, 2), sheet.title)
    output = io.BytesIO()
    workbook.save(output)
    body = output.getvalue()
    if not body or len(body) > MAX_XLSX_BYTES:
        raise ValueError("xlsx_output_limit")
    return body, observation_count


def _output_paths(settings: Settings, relative_output: str) -> tuple[Path, Path, Path]:
    if not isinstance(relative_output, str) or not relative_output.strip():
        raise ValueError("output_path_invalid")
    relative = Path(relative_output)
    if (relative.is_absolute() or relative.drive or ".." in relative.parts
            or relative.suffix.casefold() != ".xlsx"):
        raise ValueError("output_path_invalid")
    export_root = Path(settings.data_dir) / "exports"
    if not relative.is_relative_to(export_root) or relative == export_root:
        raise ValueError("output_path_invalid")
    target = settings.project_root / relative
    sidecar = Path(f"{target}.json")
    current = settings.project_root
    for part in relative.parts:
        current = current / part
        if current.is_symlink():
            raise ValueError("output_symlink")
    if not target.parent.resolve().is_relative_to(settings._inside(export_root.as_posix())):
        raise ValueError("output_path_escape")
    return target, sidecar, settings._inside(export_root.as_posix())


def _mkdir_parent(settings: Settings, target: Path, export_root: Path) -> None:
    current = settings.project_root
    for part in target.parent.relative_to(settings.project_root).parts:
        current = current / part
        if current.exists():
            if current.is_symlink() or not current.is_dir():
                raise ValueError("output_parent_invalid")
        else:
            current.mkdir()
    if not target.parent.resolve().is_relative_to(export_root):
        raise ValueError("output_path_escape")


def _publish_pair(target: Path, sidecar: Path, workbook: bytes, manifest: bytes) -> None:
    temporary_paths = []
    published = []
    try:
        for payload, suffix, final in (
            (workbook, ".xlsx.tmp", target), (manifest, ".json.tmp", sidecar),
        ):
            descriptor, name = tempfile.mkstemp(prefix=".macroatlas-xlsx-", suffix=suffix,
                                                dir=target.parent)
            temporary = Path(name)
            temporary_paths.append(temporary)
            with os.fdopen(descriptor, "wb") as stream:
                stream.write(payload)
                stream.flush()
                os.fsync(stream.fileno())
            if final.exists() or final.is_symlink():
                raise ValueError("output_already_exists")
            os.link(temporary, final)
            published.append(final)
        if os.name != "nt":
            descriptor = os.open(target.parent, os.O_RDONLY)
            try:
                os.fsync(descriptor)
            finally:
                os.close(descriptor)
    except Exception:
        for path in published:
            try:
                path.unlink()
            except OSError:
                pass
        raise
    finally:
        for path in temporary_paths:
            try:
                path.unlink()
            except OSError:
                pass


@_measured_xlsx_export("typed_xlsx")
def export_typed_series_artifact(
    settings: Settings,
    reference: Mapping,
    relative_output: str,
    *,
    freshness_policy: FreshnessPolicy = DEFAULT_FRESHNESS_POLICY,
    as_of: datetime | date | None = None,
) -> dict:
    """Revalidate and export one artifact without changing its source observations."""
    if as_of is None:
        as_of = datetime.now(UTC)
    with measure_current("parse"):
        document = load_typed_series_artifact(settings, reference)
    with measure_current("quality"):
        freshness = _reevaluate_freshness(document, freshness_policy, as_of)
    observation_count = sum(
        len(entry["contract"]["observations"]) for entry in document["series"]
    )
    input_bytes = _artifact_input_bytes(settings, reference, document)
    _record_export_shape(
        document, input_bytes=input_bytes, observation_count=observation_count,
    )
    document["_artifact_sha256"] = reference["typed_series_sha256"]
    document["_typed_series_manifest_path"] = reference["typed_series_manifest_path"]
    document["_typed_series_manifest_sha256"] = reference["typed_series_manifest_sha256"]
    with measure_current("export"):
        body, observation_count = _build_raw_workbook(document, freshness)
        target, sidecar, export_root = _output_paths(settings, relative_output)
        _mkdir_parent(settings, target, export_root)
        if target.exists() or target.is_symlink() or sidecar.exists() or sidecar.is_symlink():
            raise ValueError("output_already_exists")
        manifest_document = {
            "schema_version": "typed-series-xlsx-provenance-v1",
            "created_at": datetime.now(UTC).isoformat(),
            "output_path": target.relative_to(settings.project_root).as_posix(),
            "output_sha256": hashlib.sha256(body).hexdigest(),
            "typed_series_path": reference["typed_series_path"],
            "typed_series_sha256": reference["typed_series_sha256"],
            "typed_series_manifest_path": reference["typed_series_manifest_path"],
            "typed_series_manifest_sha256": reference["typed_series_manifest_sha256"],
            "freshness": freshness,
            "series_count": len(document["series"]),
            "observation_count": observation_count,
        }
        manifest = json.dumps(manifest_document, ensure_ascii=False, allow_nan=False,
                              sort_keys=True, indent=2).encode("utf-8")
        _publish_pair(target, sidecar, body, manifest)
    _record_export_shape(
        document,
        input_bytes=input_bytes,
        output_bytes=len(body) + len(manifest),
        observation_count=observation_count,
    )
    return {
        "status": "exported",
        "output_path": manifest_document["output_path"],
        "output_sha256": manifest_document["output_sha256"],
        "provenance_path": sidecar.relative_to(settings.project_root).as_posix(),
        "provenance_sha256": hashlib.sha256(manifest).hexdigest(),
        "series_count": manifest_document["series_count"],
        "observation_count": observation_count,
        "freshness": freshness,
    }


@_measured_xlsx_export("typed_xlsx")
def export_series_task(
    settings: Settings,
    task_id: str,
    relative_output: str,
    *,
    freshness_policy: FreshnessPolicy = DEFAULT_FRESHNESS_POLICY,
    as_of: datetime | date | None = None,
) -> dict:
    """Export the single typed artifact referenced by one completed task history."""
    with measure_current("sqlite"):
        task = HistoryStore(settings).get(task_id)
    if task.get("status") != "COMPLETED":
        raise ValueError("task_not_completed")
    references = []
    for stage in task.get("stages", []):
        result = stage.get("result") if isinstance(stage, dict) else None
        payload = result.get("payload") if isinstance(result, dict) else None
        reference = payload.get("typed_series_artifact") if isinstance(payload, dict) else None
        if isinstance(reference, Mapping):
            references.append(reference)
    if len(references) != 1:
        raise ValueError("task_must_reference_one_typed_series_artifact")
    result = export_typed_series_artifact(
        settings, references[0], relative_output,
        freshness_policy=freshness_policy, as_of=as_of,
    )
    return {"task_id": task_id, **result}


def _task_artifact_reference(settings: Settings, task_id: str) -> Mapping:
    if not isinstance(task_id, str) or not task_id.strip():
        raise ValueError("comparison_task_id_invalid")
    task = HistoryStore(settings).get(task_id)
    if task.get("status") != "COMPLETED":
        raise ValueError("comparison_task_not_completed")
    references = []
    for stage in task.get("stages", []):
        result = stage.get("result") if isinstance(stage, dict) else None
        payload = result.get("payload") if isinstance(result, dict) else None
        reference = payload.get("typed_series_artifact") if isinstance(payload, dict) else None
        if isinstance(reference, Mapping):
            references.append(reference)
    if len(references) != 1:
        raise ValueError("comparison_task_must_reference_one_typed_series_artifact")
    return references[0]


def _read_comparison_spec(settings: Settings, relative_spec: str) -> tuple[dict, bytes]:
    if not isinstance(relative_spec, str) or not relative_spec.strip():
        raise ValueError("comparison_spec_path_invalid")
    candidate = Path(relative_spec)
    if candidate.is_absolute() or candidate.drive or ".." in candidate.parts:
        raise ValueError("comparison_spec_path_invalid")
    path = settings._inside(candidate.as_posix())
    current = settings.project_root
    for part in candidate.parts:
        current = current / part
        if current.is_symlink():
            raise ValueError("comparison_spec_symlink")
    if not path.is_file() or path.stat().st_size > 64 * 1024:
        raise ValueError("comparison_spec_missing_or_too_large")
    raw = path.read_bytes()
    try:
        spec = json.loads(raw)
    except (UnicodeError, ValueError):
        raise ValueError("comparison_spec_invalid_json") from None
    if not isinstance(spec, dict) or spec.get("schema_version") != "series-comparison-spec-v1":
        raise ValueError("comparison_spec_schema_invalid")
    required = {"schema_version", "common_concept_id", "evidence", "review", "comparison", "series"}
    if set(spec) != required:
        raise ValueError("comparison_spec_fields_invalid")
    return spec, raw


def _validate_comparison_spec(spec: Mapping) -> tuple[dict, list[dict]]:
    concept_id = spec.get("common_concept_id")
    evidence = spec.get("evidence")
    review = spec.get("review")
    comparison = spec.get("comparison")
    entries = spec.get("series")
    if not isinstance(concept_id, str) or not concept_id.strip() or len(concept_id) > 200:
        raise ValueError("comparison_concept_id_invalid")
    if not isinstance(evidence, dict) or set(evidence) != {"reference", "sha256", "citation"}:
        raise ValueError("comparison_evidence_invalid")
    reference = evidence.get("reference")
    try:
        parsed = urlsplit(reference)
    except (TypeError, ValueError):
        raise ValueError("comparison_evidence_reference_invalid") from None
    if (parsed.scheme != "https" or not parsed.hostname or parsed.username or parsed.password
            or parsed.query or parsed.fragment):
        raise ValueError("comparison_evidence_reference_invalid")
    if (not isinstance(evidence.get("sha256"), str)
            or re.fullmatch(r"[0-9a-fA-F]{64}", evidence["sha256"]) is None
            or not isinstance(evidence.get("citation"), str) or not evidence["citation"].strip()):
        raise ValueError("comparison_evidence_invalid")
    if not isinstance(review, dict) or set(review) != {
            "decision", "reviewer_type", "reviewer_id", "reviewed_at"}:
        raise ValueError("comparison_review_invalid")
    if (review.get("decision") != "approved"
            or review.get("reviewer_type") not in {"AI", "human"}
            or not isinstance(review.get("reviewer_id"), str)
            or not review["reviewer_id"].strip()):
        raise ValueError("comparison_review_not_approved")
    try:
        reviewed_at = datetime.fromisoformat(review["reviewed_at"])
    except (TypeError, ValueError):
        raise ValueError("comparison_review_timestamp_invalid") from None
    if reviewed_at.tzinfo is None or reviewed_at.utcoffset() is None:
        raise ValueError("comparison_review_timestamp_invalid")
    required_comparison = {
        "date_basis": "typed_period",
        "aggregation": "none",
        "transformation": "none",
        "missing_values": "blank",
        "interpolation": "none",
        "forward_fill": "none",
        "revision_policy": "as_retrieved",
    }
    if comparison != required_comparison:
        raise ValueError("comparison_policy_unsupported")
    if not isinstance(entries, list) or not 2 <= len(entries) <= 10:
        raise ValueError("comparison_series_count_invalid")
    labels = set()
    identities = set()
    normalized = []
    for item in entries:
        if (not isinstance(item, dict) or set(item) != {"task_id", "series_id", "label"}
                or not all(isinstance(item.get(key), str) and item[key].strip()
                           for key in ("task_id", "series_id", "label"))):
            raise ValueError("comparison_series_entry_invalid")
        label = item["label"]
        if len(label) > 120 or label.casefold() in labels:
            raise ValueError("comparison_series_label_invalid_or_duplicate")
        identity = (item["task_id"], item["series_id"])
        if identity in identities:
            raise ValueError("comparison_series_identity_duplicate")
        identities.add(identity)
        labels.add(label.casefold())
        normalized.append(dict(item))
    metadata = {
        "common_concept_id": concept_id,
        "evidence_reference": evidence["reference"],
        "evidence_sha256": evidence["sha256"].lower(),
        "evidence_citation": evidence["citation"],
        "reviewer_type": review["reviewer_type"],
        "reviewer_id": review["reviewer_id"],
        "reviewed_at": reviewed_at.astimezone(UTC).isoformat(),
        **comparison,
    }
    return metadata, normalized


def _select_contract(document: Mapping, series_id: str) -> dict:
    matches = [entry for entry in document["series"]
               if entry["contract"].get("series_id") == series_id]
    if len(matches) != 1:
        raise ValueError("comparison_series_identity_not_unique")
    return matches[0]


def _check_compatibility(contracts: list[dict]) -> None:
    first = contracts[0]
    base = first["contract"]
    expected = {
        "calendar": base.get("calendar"),
        "frequency": base.get("frequency"),
        "unit": base.get("unit"),
        "currency": base.get("currency"),
        "dimension_names": sorted(base.get("dimensions", {})),
    }
    if any(value is None for value in (expected["calendar"], expected["frequency"],
                                       expected["unit"])):
        raise ValueError("comparison_metadata_unknown")
    for entry in contracts:
        contract = entry["contract"]
        observed = {
            "calendar": contract.get("calendar"),
            "frequency": contract.get("frequency"),
            "unit": contract.get("unit"),
            "currency": contract.get("currency"),
            "dimension_names": sorted(contract.get("dimensions", {})),
        }
        if observed != expected:
            raise ValueError("comparison_series_incompatible")
        periods = [item["period"] for item in contract["observations"]]
        if len(periods) != len(set(periods)):
            raise ValueError("comparison_duplicate_period")


def _add_alignment_sheet(body: bytes, contracts: list[dict], labels: list[str],
                         metadata: Mapping) -> bytes:
    workbook = load_workbook(io.BytesIO(body), data_only=False)
    sheet = workbook.create_sheet("Alignment")
    summary = (
        ("common_concept_id", metadata["common_concept_id"]),
        ("date_basis", metadata["date_basis"]),
        ("aggregation", metadata["aggregation"]),
        ("transformation", metadata["transformation"]),
        ("calendar", contracts[0]["contract"]["calendar"]),
        ("frequency", contracts[0]["contract"]["frequency"]),
        ("unit", contracts[0]["contract"]["unit"]),
        ("currency", contracts[0]["contract"]["currency"]),
        ("revision_policy", metadata["revision_policy"]),
        ("missing_values", metadata["missing_values"]),
        ("interpolation", metadata["interpolation"]),
        ("forward_fill", metadata["forward_fill"]),
    )
    for row_index, (key, value) in enumerate(summary, start=1):
        _set_literal(sheet.cell(row_index, 1), key)
        _set_literal(sheet.cell(row_index, 2), value)
    header_row = len(summary) + 2
    headers = ["typed_period"]
    for label in labels:
        headers.extend((label, f"{label}__status"))
    for column, value in enumerate(headers, start=1):
        _set_literal(sheet.cell(header_row, column), value)
    observations_by_period = []
    all_periods = set()
    for entry in contracts:
        by_period = {row["period"]: row for row in entry["contract"]["observations"]}
        observations_by_period.append(by_period)
        all_periods.update(by_period)
    for row_index, period in enumerate(sorted(all_periods), start=header_row + 1):
        _set_literal(sheet.cell(row_index, 1), period)
        column = 2
        for by_period in observations_by_period:
            observation = by_period.get(period)
            if observation is None:
                value, status = None, "no_observation"
            elif observation["status"] != "observed" or observation["raw_value"] is None:
                value, status = None, observation["status"]
            else:
                value, status = observation["raw_value"], "observed"
            _set_literal(sheet.cell(row_index, column), value)
            _set_literal(sheet.cell(row_index, column + 1), status)
            column += 2
    output = io.BytesIO()
    workbook.save(output)
    result = output.getvalue()
    if len(result) > MAX_XLSX_BYTES:
        raise ValueError("xlsx_output_limit")
    return result


@_measured_xlsx_export("series_compare_xlsx")
def export_comparison_spec(
    settings: Settings,
    spec_path: str,
    relative_output: str,
    *,
    freshness_policy: FreshnessPolicy = DEFAULT_FRESHNESS_POLICY,
    as_of: datetime | date | None = None,
) -> dict:
    """Create a comparison workbook only from explicitly reviewed compatible series."""
    if as_of is None:
        as_of = datetime.now(UTC)
    with measure_current("parse"):
        spec, spec_bytes = _read_comparison_spec(settings, spec_path)
        metadata, selectors = _validate_comparison_spec(spec)
    selected, source_documents, artifacts, freshness = [], [], [], []
    input_bytes = len(spec_bytes)
    for selector in selectors:
        with measure_current("sqlite"):
            reference = _task_artifact_reference(settings, selector["task_id"])
        with measure_current("parse"):
            document = load_typed_series_artifact(settings, reference)
        input_bytes += _artifact_input_bytes(settings, reference, document)
        entry = _select_contract(document, selector["series_id"])
        single = {"series": [entry]}
        with measure_current("quality"):
            freshness.extend(_reevaluate_freshness(single, freshness_policy, as_of))
        selected.append({**entry, "label": selector["label"], "task_id": selector["task_id"]})
        source_documents.append(document)
        artifacts.append({
            "task_id": selector["task_id"],
            "typed_series_path": reference["typed_series_path"],
            "typed_series_sha256": reference["typed_series_sha256"],
            "typed_series_manifest_path": reference["typed_series_manifest_path"],
            "typed_series_manifest_sha256": reference["typed_series_manifest_sha256"],
            "series_id": selector["series_id"],
            "label": selector["label"],
        })
    _check_compatibility(selected)
    observation_count = sum(
        len(entry["contract"]["observations"]) for entry in selected
    )
    if source_documents:
        sample_classes = {_source_sample_class(row.get("engine_id"))
                          for row in source_documents}
        sample_class = next(iter(sample_classes)) if len(sample_classes) == 1 else "unspecified"
        metrics = current_metrics()
        if metrics is not None:
            metrics.observe_sample_class(sample_class)
    _record_export_shape(
        source_documents[0] if source_documents else {},
        input_bytes=input_bytes,
        observation_count=observation_count,
        series_count=len(selected),
    )
    source_refs_by_key = {}
    for document in source_documents:
        for source_ref in document["source_refs"]:
            source_refs_by_key[(source_ref["role"], source_ref["path"])] = source_ref
    bundle = {
        "schema_version": "typed-series-artifact-v1",
        "engine_id": "series_comparison/1",
        "created_at": datetime.now(UTC).isoformat(),
        "_artifact_sha256": "multiple_inputs",
        "_typed_series_manifest_path": "multiple_inputs",
        "_typed_series_manifest_sha256": "multiple_inputs",
        "source_refs": [source_refs_by_key[key] for key in sorted(source_refs_by_key)],
        "series": selected,
        "_input_artifacts": [
            {"task_id": item["task_id"], "series_id": item["series_id"],
             "label": item["label"], "typed_series_path": item["typed_series_path"],
             "typed_series_sha256": item["typed_series_sha256"],
             "typed_series_manifest_path": item["typed_series_manifest_path"],
             "typed_series_manifest_sha256": item["typed_series_manifest_sha256"]}
            for item in artifacts
        ],
    }
    with measure_current("export"):
        body, observation_count = _build_raw_workbook(
            bundle, freshness, comparison_metadata=metadata,
        )
        body = _add_alignment_sheet(body, selected, [item["label"] for item in selected], metadata)
        spec_hash = hashlib.sha256(spec_bytes).hexdigest()
        target, sidecar, export_root = _output_paths(settings, relative_output)
        _mkdir_parent(settings, target, export_root)
        if target.exists() or target.is_symlink() or sidecar.exists() or sidecar.is_symlink():
            raise ValueError("output_already_exists")
        manifest_document = {
            "schema_version": "series-comparison-xlsx-provenance-v1",
            "created_at": datetime.now(UTC).isoformat(),
            "output_path": target.relative_to(settings.project_root).as_posix(),
            "output_sha256": hashlib.sha256(body).hexdigest(),
            "comparison_spec_path": spec_path,
            "comparison_spec_sha256": spec_hash,
            "comparison_review": metadata,
            "inputs": artifacts,
            "freshness": freshness,
            "series_count": len(selected),
            "observation_count": observation_count,
            "note": "Evidence URL and reviewer identity are recorded but not authenticated by this tool.",
        }
        manifest = json.dumps(manifest_document, ensure_ascii=False, allow_nan=False,
                              sort_keys=True, indent=2).encode("utf-8")
        _publish_pair(target, sidecar, body, manifest)
    _record_export_shape(
        source_documents[0] if source_documents else {},
        input_bytes=input_bytes,
        output_bytes=len(body) + len(manifest),
        observation_count=observation_count,
        series_count=len(selected),
    )
    return {
        "status": "exported",
        "output_path": manifest_document["output_path"],
        "output_sha256": manifest_document["output_sha256"],
        "provenance_path": sidecar.relative_to(settings.project_root).as_posix(),
        "provenance_sha256": hashlib.sha256(manifest).hexdigest(),
        "comparison_spec_sha256": spec_hash,
        "series_count": len(selected),
        "observation_count": observation_count,
        "freshness": freshness,
        "reviewer_type": metadata["reviewer_type"],
        "evidence_identity_authenticated": False,
    }
